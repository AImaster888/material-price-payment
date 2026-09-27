# -*- coding: utf-8 -*-
"""
verify_payment.py — 估驗計價逐項核算引擎。

適用對象：台灣公共工程「估驗計價」文件常見版型——一張表裡橫向分成幾個欄位群組
（常見：原契約 / 前期累計 / 本期完成 / 累計至本期），每個群組底下再分「數量」「複價」
（原契約群組通常還有「單位」「單價」）。這是政府機關常用的標準格式，不同機關、
不同案子的欄位「字母位置」常常不一樣，所以本引擎用**搜尋標題文字**來定位欄位，
不假設固定的欄位字母。

核心檢查（對應實際審查估驗計價文件時最容易出錯、也最容易被忽略的地方）：

  1. extended-price   單價 x 數量 = 複價（本期完成群組，逐項）
  2. within-period     前期累計 + 本期完成 = 累計至本期（數量、複價都查）
  3. cross-period      A檔「累計至本期」= B檔「前期累計」（前後兩期銜接，逐項配對）
  4. subitem-sum       有子項的小計列，其複價應等於底下子項複價加總（自動偵測 SUM 公式範圍）
  5. display-text      不管底層數值，只比對 Excel 實際「畫面顯示文字」是否一致
                       （見 render_excel.py 開頭說明——這是本工具存在的主要原因）

用法：
    python verify_payment.py <前一期.xlsx> <本期.xlsx> --sheet 請款明細表

    只給一份檔案時，只會做 extended-price / within-period / subitem-sum 三項
    （cross-period 需要兩份檔案才能比對銜接：前一期「累計至本期」應等於本期「前期累計」，
    所以順序一定是「前一期在前、本期在後」，不要顛倒）。

exit code：
    0 = 已跑的檢查全數相符
    1 = 有異常
    2 = 有檢查因欄位不齊被略過（**不可當成「全部相符」**）
    3 = 檔案不存在／分頁名稱錯／版型偵測失敗，根本沒跑成

設計原則：AI 不重算任何數字——所有數字都是這支腳本用 openpyxl 讀出來、用 Python
算出來的，Claude 只負責讀腳本輸出、翻譯成人話、決定下一步。
"""
import argparse
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).parent))
from render_excel import render_display  # noqa: E402

__version__ = '1.3.0'  # 2026-09-27

GROUP_NAMES = ['原契約', '原派工', '第一次變更設計', '第一次變更設計(議價前)', '前期累計', '本期完成', '累計至本期']
GROUP_SYNONYMS = {'原派工': '原契約'}  # 不同期的檔案可能對同一個群組用不同標題文字
SUBCOL_NAMES = ['單位', '數量', '單價', '複價']
CODE_HEADER_CANDIDATES = ['項目', '項次']
NAME_HEADER_CANDIDATES = ['工程項目', '項目名稱', '名稱']


class LayoutError(Exception):
    pass


class InputError(Exception):
    """檔案不存在、分頁名稱打錯這類「使用者輸入問題」，跟核算結果無關。"""
    pass


class SheetLayout:
    def __init__(self, ws):
        self.ws = ws
        self.code_col = None
        self.name_col = None
        self.group_header_row = None
        self.subheader_row = None
        self.groups = {}  # group_name -> {subcol_name: col_idx}
        self.first_data_row = None
        self.last_data_row = None
        self.scan_stop_reason = None  # 資料列掃描為什麼停在 last_data_row（一定要讓使用者看到）
        self._detect()

    def _find_group_header_row(self):
        max_scan_row = min(15, self.ws.max_row)
        max_scan_col = min(60, self.ws.max_column)
        best_row, best_hits = None, 0
        for r in range(1, max_scan_row + 1):
            hits = 0
            for c in range(1, max_scan_col + 1):
                v = self.ws.cell(row=r, column=c).value
                if isinstance(v, str) and v.strip() in GROUP_NAMES:
                    hits += 1
            if hits > best_hits:
                best_hits, best_row = hits, r
        if best_hits == 0:
            raise LayoutError('找不到任何群組標題（原契約/前期累計/本期完成/累計至本期），無法自動偵測欄位版型')
        return best_row

    def _detect(self):
        ws = self.ws
        self.group_header_row = self._find_group_header_row()
        self.subheader_row = self.group_header_row + 1
        max_col = min(80, ws.max_column)

        # 找每個群組標題出現的欄位範圍
        group_starts = []  # (col, name)
        for c in range(1, max_col + 1):
            v = ws.cell(row=self.group_header_row, column=c).value
            if isinstance(v, str) and v.strip() in GROUP_NAMES:
                name = v.strip()
                name = GROUP_SYNONYMS.get(name, name)
                group_starts.append((c, name))
        group_starts.append((max_col + 1, None))  # sentinel

        for i in range(len(group_starts) - 1):
            start_col, name = group_starts[i]
            end_col = group_starts[i + 1][0]
            cols = {}
            for c in range(start_col, end_col):
                v = ws.cell(row=self.subheader_row, column=c).value
                if isinstance(v, str) and v.strip() in SUBCOL_NAMES:
                    cols[v.strip()] = c
            self.groups[name] = cols

        # 找項次代碼欄、工程項目名稱欄：在群組標題列往前找 "項目"/"項次"/"工程項目"
        for c in range(1, max_col + 1):
            v = ws.cell(row=self.group_header_row, column=c).value
            if isinstance(v, str):
                vs = v.strip()
                if self.code_col is None and vs in CODE_HEADER_CANDIDATES:
                    self.code_col = c
                if self.name_col is None and vs in NAME_HEADER_CANDIDATES:
                    self.name_col = c
        if self.code_col is None:
            self.code_col = 1
        if self.name_col is None:
            self.name_col = self.code_col + 1

        # 資料列範圍：從 subheader_row+1 開始，直到代碼欄與名稱欄連續多列都是空的。
        # 表格下方常接著簽名欄／日期註記等文字列（例如「施工承攬廠商（編製）：」），
        # 這些列的代碼/名稱欄不是空的，但數量／複價欄不該出現文字——用這點來排除，
        # 避免把表尾的簽名區也當成資料列。
        numeric_cols = []
        for cols in self.groups.values():
            for key in ('數量', '複價'):
                if key in cols:
                    numeric_cols.append(cols[key])

        data_start = self.subheader_row + 1
        last_row = data_start
        empty_streak = 0
        reason = '掃到工作表最後一列（第 {0} 列）'.format(ws.max_row)
        for r in range(data_start, ws.max_row + 1):
            a = ws.cell(row=r, column=self.code_col).value
            b = ws.cell(row=r, column=self.name_col).value
            if a is None and b is None:
                empty_streak += 1
                if empty_streak >= 8:
                    reason = '第 {0} 列起連續 8 列空白'.format(r - 7)
                    break
                continue
            # ws 是公式模式的工作表，數量/複價欄合法值不是 None 就是公式字串（"=..."）；
            # 如果出現不是公式的純文字，就是表尾註記誤入資料範圍
            has_stray_text = any(
                isinstance(ws.cell(row=r, column=c).value, str)
                and not ws.cell(row=r, column=c).value.strip().startswith('=')
                for c in numeric_cols
            )
            if has_stray_text:
                reason = '第 {0} 列的數量/複價欄出現非公式文字（研判是表尾註記/簽名欄）'.format(r)
                break
            empty_streak = 0
            last_row = r
            # 表格慣例上「總計」是資料區的最後一列，之後接的是簽名欄／日期註記，
            # 遇到就停止往下掃，避免把表尾雜訊也當成資料列。
            # 但「本頁小計」「承前頁總計」這種分頁彙總列是資料區中途的列，不能當結尾。
            if isinstance(b, str) and '總計' in b and not any(
                    k in b for k in ('本頁', '次頁', '承前', '接次')):
                reason = '第 {0} 列名稱含「總計」，視為資料區最後一列'.format(r)
                break
        self.first_data_row = data_start
        self.last_data_row = last_row
        self.scan_stop_reason = reason

    def col(self, group, subcol):
        g = self.groups.get(group)
        if not g:
            return None
        return g.get(subcol)

    def describe(self):
        lines = ['偵測到的欄位版型：',
                 '  項次欄={0}, 名稱欄={1}, 群組標題列={2}, 子標題列={3}, 資料列={4}~{5}'.format(
                     self.code_col, self.name_col, self.group_header_row, self.subheader_row,
                     self.first_data_row, self.last_data_row)]
        for name, cols in self.groups.items():
            if cols:
                lines.append('  【{0}】 {1}'.format(name, cols))
        lines.append('  掃描停止原因：{0}'.format(self.scan_stop_reason))
        lines.append('  ↑ 如果表格實際的資料列不只到第 {0} 列，代表掃描提早停了，'
                     '結果不完整，要回報使用者。'.format(self.last_data_row))
        return '\n'.join(lines)


def _num(v):
    """把讀到的值收斂成數字或 None，防止表尾雜訊文字混進算式。"""
    return v if isinstance(v, (int, float)) else None


def _round_half_up(v):
    """四捨五入到整數（Excel ROUND 的行為）。

    不能用 Python 內建的 round()——它是「銀行家捨入」(round-half-to-even)：
    round(1234.5) 會給 1234，但 Excel 的 ROUND(1234.5,0) 給 1235，
    單價帶小數時就會憑空生出假異常。
    """
    return float(Decimal(repr(float(v))).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def _key(v):
    """項次代碼／名稱拿來配對用的正規化鍵。

    兩份檔案同一個項次，一份可能存成數字 1、另一份存成文字 "1"，
    或是縮排空格、全形空格不一樣——不正規化就會誤判成「找不到對應項次」。
    """
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r'\s+', '', str(v))
    return s or None


def check_extended_price(wf, wv, layout, group='本期完成', tol=0.5):
    """單價 x 數量 = 複價。單價優先找同群組內的，找不到就退回原契約群組的單價。"""
    qty_col = layout.col(group, '數量')
    amt_col = layout.col(group, '複價')
    price_col = layout.col(group, '單價') or layout.col('原契約', '單價')
    results = []
    if qty_col is None or amt_col is None or price_col is None:
        return results, '缺少必要欄位（單價/數量/複價），略過此檢查'
    for r in range(layout.first_data_row, layout.last_data_row + 1):
        code = wv.cell(row=r, column=layout.code_col).value
        name = wv.cell(row=r, column=layout.name_col).value
        if code is None and name is None:
            continue
        price = wv.cell(row=r, column=price_col).value
        qty = wv.cell(row=r, column=qty_col).value
        amt = wv.cell(row=r, column=amt_col).value
        amt_formula = wf.cell(row=r, column=amt_col).value
        if price is None or qty is None:
            continue  # 小計/特殊列，非逐項乘積
        expected = _round_half_up(price * qty)
        ok = amt is not None and abs(expected - amt) < tol
        results.append(dict(row=r, code=code, name=name, price=price, qty=qty,
                             expected=expected, actual=amt, ok=ok, formula=amt_formula))
    return results, None


def check_within_period(wv, layout, tol=0.5):
    """前期累計 + 本期完成 = 累計至本期（數量、複價都查）。"""
    results = []
    prev_q = layout.col('前期累計', '數量')
    prev_a = layout.col('前期累計', '複價')
    cur_q = layout.col('本期完成', '數量')
    cur_a = layout.col('本期完成', '複價')
    tot_q = layout.col('累計至本期', '數量')
    tot_a = layout.col('累計至本期', '複價')
    if None in (prev_a, cur_a, tot_a):
        return results, '缺少前期/本期/累計複價欄位，略過此檢查'
    for r in range(layout.first_data_row, layout.last_data_row + 1):
        code = wv.cell(row=r, column=layout.code_col).value
        name = wv.cell(row=r, column=layout.name_col).value
        if code is None and name is None:
            continue
        pq = _num(wv.cell(row=r, column=prev_q).value) if prev_q else None
        pa = _num(wv.cell(row=r, column=prev_a).value)
        cq = _num(wv.cell(row=r, column=cur_q).value) if cur_q else None
        ca = _num(wv.cell(row=r, column=cur_a).value)
        tq = _num(wv.cell(row=r, column=tot_q).value) if tot_q else None
        ta = _num(wv.cell(row=r, column=tot_a).value)
        if pq is None and cq is None and tq is None:
            # 這一列本來就沒有「數量」概念（常見於「式」計價的大項小計列），
            # 不是漏填，不當異常處理
            exp_q, q_ok = None, True
        else:
            exp_q = (pq or 0) + (cq or 0)
            q_ok = tq is not None and abs(exp_q - tq) < 0.0001
        if pa is None and ca is None and ta is None:
            # 同上：這一列三個複價欄都是空的（大項標題列、表尾註記列），
            # 不是「累計欄漏填」，不當異常處理
            exp_a, a_ok = None, True
        else:
            exp_a = (pa or 0) + (ca or 0)
            a_ok = ta is not None and abs(exp_a - ta) < tol
        results.append(dict(row=r, code=code, name=name,
                             prev_qty=pq, prev_amt=pa, cur_qty=cq, cur_amt=ca,
                             tot_qty=tq, tot_amt=ta, exp_qty=exp_q, exp_amt=exp_a,
                             qty_ok=q_ok, amt_ok=a_ok))
    return results, None


def _build_b_index(wv_b, layout_b):
    """建 B 檔的配對索引：(依代碼, 依名稱, 重複鍵集合)。鍵都經過 `_key()` 正規化。

    同一個鍵出現兩列以上時不能亂配——記進 dup，配對時直接判為「無法配對」，
    不要靜靜地挑其中一列來比。
    """
    by_code, by_name, dup = {}, {}, set()
    for r in range(layout_b.first_data_row, layout_b.last_data_row + 1):
        code = _key(wv_b.cell(row=r, column=layout_b.code_col).value)
        # 沒有項次代碼的列（常見於「總計」「物價調整金額」這類彙總列），改用名稱比對
        name = _key(wv_b.cell(row=r, column=layout_b.name_col).value)
        target, key = (by_code, code) if code else (by_name, name)
        if key is None:
            continue
        if key in target:
            dup.add(key)
        else:
            target[key] = r
    return by_code, by_name, dup


def check_cross_period(wv_a, layout_a, wv_b, layout_b, tol_qty=0.0001, tol_amt=0.5):
    """A檔「累計至本期」應等於 B檔「前期累計」，用項次代碼配對逐列比對。"""
    a_q = layout_a.col('累計至本期', '數量')
    a_a = layout_a.col('累計至本期', '複價')
    b_q = layout_b.col('前期累計', '數量')
    b_a = layout_b.col('前期累計', '複價')
    if a_a is None or b_a is None:
        return [], 'A檔缺累計至本期複價欄，或B檔缺前期累計複價欄，略過此檢查'

    b_index_by_code, b_index_by_name, b_dup = _build_b_index(wv_b, layout_b)

    results = []
    used_b_rows = set()
    for r in range(layout_a.first_data_row, layout_a.last_data_row + 1):
        code = wv_a.cell(row=r, column=layout_a.code_col).value
        name = wv_a.cell(row=r, column=layout_a.name_col).value
        if code is None and name is None:
            continue
        aq = wv_a.cell(row=r, column=a_q).value if a_q else None
        aa = wv_a.cell(row=r, column=a_a).value
        k_code, k_name = _key(code), _key(name)
        key = k_code or k_name
        if key in b_dup:
            results.append(dict(a_row=r, b_row=None, code=code, name=name,
                                 a_qty=aq, a_amt=aa, b_qty=None, b_amt=None,
                                 qty_ok=False, amt_ok=False, matched=False,
                                 note='本期有多列的項次/名稱都是「{0}」，無法確定要跟哪一列比'.format(key)))
            continue
        br = b_index_by_code.get(k_code) if k_code else b_index_by_name.get(k_name)
        if br is None:
            results.append(dict(a_row=r, b_row=None, code=code, name=name,
                                 a_qty=aq, a_amt=aa, b_qty=None, b_amt=None,
                                 qty_ok=False, amt_ok=False, matched=False,
                                 note='本期找不到對應項次/名稱'))
            continue
        used_b_rows.add(br)
        bq = wv_b.cell(row=br, column=b_q).value if b_q else None
        ba = wv_b.cell(row=br, column=b_a).value
        qty_ok = (a_q is None or b_q is None) or (abs((aq or 0) - (bq or 0)) < tol_qty)
        amt_ok = abs((aa or 0) - (ba or 0)) < tol_amt
        results.append(dict(a_row=r, b_row=br, code=code, name=name,
                             a_qty=aq, a_amt=aa, b_qty=bq, b_amt=ba,
                             qty_ok=qty_ok, amt_ok=amt_ok, matched=True, note=''))

    # 反向再掃一次：本期有、前一期沒有的列。本期新增的項目「前期累計」本來就該是 0；
    # 若它的前期累計有金額，卻在前一期表上找不到這一項，就是銜接對不起來。
    for r in range(layout_b.first_data_row, layout_b.last_data_row + 1):
        if r in used_b_rows:
            continue
        ba = _num(wv_b.cell(row=r, column=b_a).value)
        if not ba:
            continue
        results.append(dict(a_row=None, b_row=r,
                             code=wv_b.cell(row=r, column=layout_b.code_col).value,
                             name=wv_b.cell(row=r, column=layout_b.name_col).value,
                             a_qty=None, a_amt=None,
                             b_qty=wv_b.cell(row=r, column=b_q).value if b_q else None,
                             b_amt=ba, qty_ok=False, amt_ok=False, matched=False,
                             note='前一期找不到這一項，但本期的前期累計有金額'))
    return results, None


def check_subitem_sum(wf, wv, layout, group='本期完成'):
    """偵測小計列的 SUM(x:y) 公式，驗算子項加總是否等於小計。"""
    amt_col = layout.col(group, '複價')
    if amt_col is None:
        return [], '缺少該群組複價欄，略過此檢查'
    results = []
    # 允許 $ 絕對位址（=SUM($J$7:$J$8)）與小寫 sum；範圍必須落在同一欄
    pattern = re.compile(r'^=SUM\(\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)\)$', re.IGNORECASE)
    for r in range(layout.first_data_row, layout.last_data_row + 1):
        f = wf.cell(row=r, column=amt_col).value
        if not isinstance(f, str):
            continue
        m = pattern.match(f.replace(' ', ''))
        if not m:
            continue
        col_from, col_to = m.group(1).upper(), m.group(3).upper()
        if col_from != col_to:
            continue  # 跨欄範圍（=SUM(A1:B5)）不是小計列的加總，不處理
        child_col = column_index_from_string(col_from)
        a, b = int(m.group(2)), int(m.group(4))
        child_sum = 0
        for cr in range(a, b + 1):
            child_sum += _num(wv.cell(row=cr, column=child_col).value) or 0
        actual = wv.cell(row=r, column=amt_col).value or 0
        code = wv.cell(row=r, column=layout.code_col).value
        name = wv.cell(row=r, column=layout.name_col).value
        ok = abs(child_sum - actual) < 0.5
        results.append(dict(row=r, code=code, name=name, range='{0}:{1}'.format(a, b),
                             child_sum=child_sum, actual=actual, ok=ok))
    return results, None


DISPLAY_CAUSE_FORMAT = '格式不同（數字相同）'
DISPLAY_CAUSE_VALUE = '數字不同'


def _same_display(d1, d2):
    """「-」（0 用會計格式顯示）跟空白格，紙本上都是「沒有數字」，不算畫面不一致。"""
    empty = ('', '-')
    return d1 == d2 or (d1 in empty and d2 in empty)


def _same_value(v1, v2, tol):
    return abs((_num(v1) or 0) - (_num(v2) or 0)) < tol


def check_display_text_cross(wv_a, wf_a, layout_a, wv_b, wf_b, layout_b, tol_qty=0.0001, tol_amt=0.5):
    """比對 A檔「累計至本期」vs B檔「前期累計」的畫面顯示文字（不論底層數值）。

    顯示不一致時，另外比對底層數值標出 cause：格式問題（統一格式即可）跟數字真的
    不同（要改金額）處理方式完全不同，不能讓讀報告的人或 AI 自己去猜。
    容錯值跟 check_cross_period 一致。
    """
    a_q, a_a = layout_a.col('累計至本期', '數量'), layout_a.col('累計至本期', '複價')
    b_q, b_a = layout_b.col('前期累計', '數量'), layout_b.col('前期累計', '複價')
    if a_a is None or b_a is None:
        return [], '缺少必要欄位，略過此檢查'
    b_index_by_code, b_index_by_name, b_dup = _build_b_index(wv_b, layout_b)
    results = []
    for r in range(layout_a.first_data_row, layout_a.last_data_row + 1):
        code = wv_a.cell(row=r, column=layout_a.code_col).value
        name = wv_a.cell(row=r, column=layout_a.name_col).value
        if code is None and name is None:
            continue
        k_code, k_name = _key(code), _key(name)
        if (k_code or k_name) in b_dup:
            continue  # 配不出唯一對象，交給 cross_period 去報「無法配對」
        br = b_index_by_code.get(k_code) if k_code else b_index_by_name.get(k_name)
        if br is None:
            continue
        qd1 = qd2 = None
        q_same = True
        if a_q and b_q:
            qv1, qv2 = wv_a.cell(row=r, column=a_q).value, wv_b.cell(row=br, column=b_q).value
            qd1 = render_display(qv1, wf_a.cell(row=r, column=a_q).number_format)
            qd2 = render_display(qv2, wf_b.cell(row=br, column=b_q).number_format)
            q_same = _same_value(qv1, qv2, tol_qty)
        av1, av2 = wv_a.cell(row=r, column=a_a).value, wv_b.cell(row=br, column=b_a).value
        ad1 = render_display(av1, wf_a.cell(row=r, column=a_a).number_format)
        ad2 = render_display(av2, wf_b.cell(row=br, column=b_a).number_format)
        q_match = (qd1 is None) or _same_display(qd1, qd2)
        a_match = _same_display(ad1, ad2)
        cause = ''
        if not (q_match and a_match):
            value_differs = (not q_match and not q_same) or (not a_match and not _same_value(av1, av2, tol_amt))
            cause = DISPLAY_CAUSE_VALUE if value_differs else DISPLAY_CAUSE_FORMAT
        results.append(dict(row=r, code=code, name=name, cause=cause,
                             qty_disp_a=qd1, qty_disp_b=qd2, qty_match=q_match,
                             amt_disp_a=ad1, amt_disp_b=ad2, amt_match=a_match))
    return results, None


def _load(path, sheet_name):
    """讀出 (公式版工作表, 值版工作表)。檔案/分頁有問題就丟 InputError，不要讓它變 traceback。"""
    if not Path(path).exists():
        raise InputError('找不到檔案：{0}'.format(path))
    wb_f = openpyxl.load_workbook(path, data_only=False)
    if sheet_name not in wb_f.sheetnames:
        raise InputError('檔案「{0}」裡沒有分頁「{1}」。這個檔案的分頁有：{2}'.format(
            Path(path).name, sheet_name, '、'.join(wb_f.sheetnames)))
    wb_v = openpyxl.load_workbook(path, data_only=True)
    return wb_f[sheet_name], wb_v[sheet_name]


def run_all(path_a, sheet_name, path_b=None, price_group='本期完成'):
    wf_a, wv_a = _load(path_a, sheet_name)
    layout_a = SheetLayout(wf_a)

    out = dict(file_a=str(path_a), sheet=sheet_name, layout_a=layout_a.describe())
    out['extended_price'], out['extended_price_skip'] = check_extended_price(wf_a, wv_a, layout_a, price_group)
    out['within_period'], out['within_period_skip'] = check_within_period(wv_a, layout_a)
    out['subitem_sum'], out['subitem_sum_skip'] = check_subitem_sum(wf_a, wv_a, layout_a, price_group)

    if path_b:
        wf_b, wv_b = _load(path_b, sheet_name)
        layout_b = SheetLayout(wf_b)
        out['file_b'] = str(path_b)
        out['layout_b'] = layout_b.describe()
        # 本期自己也要核單檔三項——審查的重點本來就是本期，只核前一期等於漏掉本期的錯
        out['extended_price_b'], out['extended_price_b_skip'] = check_extended_price(wf_b, wv_b, layout_b, price_group)
        out['within_period_b'], out['within_period_b_skip'] = check_within_period(wv_b, layout_b)
        out['subitem_sum_b'], out['subitem_sum_b_skip'] = check_subitem_sum(wf_b, wv_b, layout_b, price_group)
        out['cross_period'], out['cross_period_skip'] = check_cross_period(wv_a, layout_a, wv_b, layout_b)
        out['display_text'], out['display_text_skip'] = check_display_text_cross(
            wv_a, wf_a, layout_a, wv_b, wf_b, layout_b)
    return out


def print_summary(out):
    print('【前一期／本檔】')
    print(out['layout_a'])
    if 'layout_b' in out:
        print()
        print('【本期】')
        print(out['layout_b'])
    print()
    print('=== 核算結果 ===')
    for key in CHECK_KEYS:
        if key not in out:
            continue
        title = check_title(key, out)
        skip = out.get(key + '_skip')
        if skip:
            print('- {0}：⚠ 略過未檢查 — {1}'.format(title, skip))
            continue
        bad = [x for x in out[key] if not _row_ok(key, x)]
        print('- {0}：共 {1} 項，異常 {2} 項{3}'.format(title, len(out[key]), len(bad), _cause_breakdown(key, bad)))
        for x in bad[:20]:
            print('    row/code={0} {1} -> {2}'.format(x.get('row', x.get('a_row')), x.get('code'), x))
        if len(bad) > 20:
            print('    …另有 {0} 筆異常沒列出來，跑 --md 看完整報告'.format(len(bad) - 20))
    print()
    print(conclusion_line(out))


def _cause_breakdown(key, bad):
    if key != 'display_text' or not bad:
        return ''
    n_fmt = sum(1 for x in bad if x['cause'] == DISPLAY_CAUSE_FORMAT)
    return '，其中「{0}」{1} 項（統一格式即可）、「{2}」{3} 項（金額或數量要修正）'.format(
        DISPLAY_CAUSE_FORMAT, n_fmt, DISPLAY_CAUSE_VALUE, len(bad) - n_fmt)


def summarize(out):
    """回傳 (異常筆數, 被略過的檢查標題清單)。exit code 跟結論都由這裡決定。"""
    bad = 0
    skipped = []
    for key in CHECK_KEYS:
        if key not in out:
            continue
        if out.get(key + '_skip'):
            skipped.append(check_title(key, out))
            continue
        bad += sum(1 for x in out[key] if not _row_ok(key, x))
    return bad, skipped


def conclusion_line(out):
    bad, skipped = summarize(out)
    text = ('結論：共發現 {0} 筆異常。'.format(bad) if bad
            else '結論：已跑的檢查全數相符，未發現異常。')
    if skipped:
        text += '\n⚠ 但有 {0} 項檢查因為欄位不齊沒跑到（{1}），不能回報「全部相符」。'.format(
            len(skipped), '、'.join(skipped))
    return text


def _base(key):
    """'within_period_b'（本期那份的單檔檢查）跟 'within_period' 用同一套判定與表格。"""
    return key[:-2] if key.endswith('_b') else key


def check_title(key, out):
    """兩期模式下單檔檢查會跑兩次，標題要分得出是哪一期；單檔模式不加期別。"""
    title = CHECK_TITLES[_base(key)][0]
    if key.endswith('_b'):
        return title + '（本期）'
    if 'file_b' in out and key in SINGLE_FILE_KEYS:
        return title + '（前一期）'
    return title


def _row_ok(key, x):
    key = _base(key)
    if key == 'extended_price':
        return x['ok']
    if key == 'within_period':
        return x['qty_ok'] and x['amt_ok']
    if key == 'subitem_sum':
        return x['ok']
    if key == 'cross_period':
        return x['matched'] and x['qty_ok'] and x['amt_ok']
    if key == 'display_text':
        return x['qty_match'] and x['amt_match']
    return True


SINGLE_FILE_KEYS = ('extended_price', 'within_period', 'subitem_sum')
CHECK_KEYS = SINGLE_FILE_KEYS + tuple(k + '_b' for k in SINGLE_FILE_KEYS) + ('cross_period', 'display_text')

CHECK_TITLES = {
    'extended_price': ('單價 x 本期數量 = 複價', '逐項核對每一列的單價乘上本期數量，是否等於複價欄位的實際值'),
    'within_period': ('前期累計 + 本期完成 = 累計至本期', '確認前後兩欄相加是否等於「累計至本期」欄位（同一份檔案內部一致性）'),
    'subitem_sum': ('子項加總反查', '偵測 SUM 公式的小計列，確認底下子項複價加總是否等於小計'),
    'cross_period': ('跨期銜接比對', '前一期「累計至本期」應等於本期「前期累計」，用項次代碼（或名稱）逐列配對'),
    'display_text': ('畫面顯示文字比對', '不論底層數值是否相等，只比對 Excel 依各自儲存格格式顯示出來的文字是否一致'),
}


def _fmt_num(v):
    if v is None:
        return '—'
    if isinstance(v, float):
        return ('{0:,.6f}'.format(v)).rstrip('0').rstrip('.') or '0'
    return '{0:,}'.format(v)


def build_markdown(result):
    lines = []
    lines.append('# 估驗計價逐項核算報告')
    lines.append('')
    lines.append('- 分頁：`{0}`'.format(result['sheet']))
    lines.append('- 前一期檔案：`{0}`'.format(result['file_a']))
    if 'file_b' in result:
        lines.append('- 本期檔案：`{0}`'.format(result['file_b']))
    lines.append('')
    lines.append('```')
    lines.append(result['layout_a'])
    if 'layout_b' in result:
        lines.append('')
        lines.append(result['layout_b'])
    lines.append('```')
    lines.append('')

    for key in CHECK_KEYS:
        if key not in result:
            continue
        title, desc = check_title(key, result), CHECK_TITLES[_base(key)][1]
        skip = result.get(key + '_skip')
        rows = result[key]
        lines.append('## {0}'.format(title))
        lines.append('')
        lines.append(desc)
        lines.append('')
        if skip:
            lines.append('> ⚠ 略過：{0}'.format(skip))
            lines.append('')
            continue
        bad = [x for x in rows if not _row_ok(key, x)]
        lines.append('共 {0} 項，異常 {1} 項{2}'.format(
            len(rows), len(bad), '（全數正確 ✓）' if not bad else _cause_breakdown(key, bad)))
        lines.append('')
        if bad:
            base = _base(key)
            if base == 'extended_price':
                lines.append('| 項次 | 工程項目 | 單價 | 數量 | 應為複價 | 實際複價 |')
                lines.append('|---|---|---:|---:|---:|---:|')
                for x in bad:
                    lines.append('| {0} | {1} | {2} | {3} | {4} | {5} |'.format(
                        (x['code'] or '—'), x['name'], _fmt_num(x['price']), _fmt_num(x['qty']),
                        _fmt_num(x['expected']), _fmt_num(x['actual'])))
            elif base == 'within_period':
                lines.append('| 項次 | 工程項目 | 前期複價 | 本期複價 | 應為累計 | 實際累計 | 問題 |')
                lines.append('|---|---|---:|---:|---:|---:|---|')
                for x in bad:
                    problem = []
                    if not x['qty_ok']:
                        problem.append('數量')
                    if not x['amt_ok']:
                        problem.append('複價')
                    lines.append('| {0} | {1} | {2} | {3} | {4} | {5} | {6} |'.format(
                        (x['code'] or '—'), x['name'], _fmt_num(x['prev_amt']), _fmt_num(x['cur_amt']),
                        _fmt_num(x['exp_amt']), _fmt_num(x['tot_amt']), '、'.join(problem)))
            elif base == 'subitem_sum':
                lines.append('| 小計列 | 工程項目 | 子項範圍 | 子項加總 | 小計實際值 |')
                lines.append('|---|---|---|---:|---:|')
                for x in bad:
                    lines.append('| {0} | {1} | {2} | {3} | {4} |'.format(
                        (x['code'] or '—'), x['name'], x['range'], _fmt_num(x['child_sum']), _fmt_num(x['actual'])))
            elif base == 'cross_period':
                lines.append('| 項次 | 工程項目 | 前一期累計數量 | 本期前期數量 | 前一期累計複價 | 本期前期複價 | 問題 |')
                lines.append('|---|---|---:|---:|---:|---:|---|')
                for x in bad:
                    if not x['matched']:
                        problem = x.get('note') or '本期找不到對應項次/名稱'
                    else:
                        problem = '、'.join(
                            (['數量'] if not x['qty_ok'] else []) + (['複價'] if not x['amt_ok'] else []))
                    lines.append('| {0} | {1} | {2} | {3} | {4} | {5} | {6} |'.format(
                        (x['code'] or '—'), x['name'], _fmt_num(x['a_qty']), _fmt_num(x['b_qty']),
                        _fmt_num(x['a_amt']), _fmt_num(x['b_amt']), problem))
            elif base == 'display_text':
                lines.append('| 項次 | 工程項目 | 前一期顯示(數量) | 本期顯示(數量) | 前一期顯示(複價) | 本期顯示(複價) | 原因 |')
                lines.append('|---|---|---|---|---|---|---|')
                for x in bad:
                    lines.append('| {0} | {1} | {2} | {3} | {4} | {5} | {6} |'.format(
                        (x['code'] or '—'), x['name'], x['qty_disp_a'], x['qty_disp_b'],
                        x['amt_disp_a'], x['amt_disp_b'], x['cause']))
        lines.append('')

    lines.append('---')
    lines.append('**{0}**'.format(conclusion_line(result).replace('\n', '**\n\n**')))
    return '\n'.join(lines)


def md_report_path(period_a, sheet_name):
    """檔名要帶分頁名稱：同一組檔案通常兩張表都要核，不帶的話後跑的會蓋掉先跑的。"""
    p = Path(period_a)
    return p.with_name('{0}-{1}-核算報告.md'.format(p.stem, sheet_name))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='估驗計價逐項核算引擎')
    ap.add_argument('period_a', help='前一期（或唯一一份）估驗計價 xlsx 路徑')
    ap.add_argument('period_b', nargs='?', default=None, help='本期估驗計價 xlsx 路徑（選填，給了才做跨期比對，順序不可顛倒）')
    ap.add_argument('--sheet', required=True, help='要核算的分頁名稱，例如「請款明細表」')
    ap.add_argument('--price-group', default='本期完成', help='複價核算要用哪個群組（預設：本期完成）')
    ap.add_argument('--md', action='store_true', help='額外輸出 <period_a檔名>-<分頁名稱>-核算報告.md（跟 period_a 同目錄）')
    ap.add_argument('--version', action='version', version=__version__)
    args = ap.parse_args()

    try:
        result = run_all(args.period_a, args.sheet, args.period_b, args.price_group)
    except (InputError, LayoutError) as e:
        # 檔案/分頁/版型的問題不是「核算發現異常」，exit code 要分得開
        print('無法核算：{0}'.format(e))
        sys.exit(3)

    print_summary(result)

    if args.md:
        md_path = md_report_path(args.period_a, args.sheet)
        md_path.write_text(build_markdown(result), encoding='utf-8')
        print('\n已輸出：{0}'.format(md_path))

    bad, skipped = summarize(result)
    sys.exit(1 if bad else (2 if skipped else 0))
