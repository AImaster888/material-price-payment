# -*- coding: utf-8 -*-
"""
產生「練習用」兩期估驗計價範例檔（全部是假資料，可公開）。

故意埋了真實案件最常見的 4 種錯，給使用手冊當範例、也給新電腦裝完技能後當驗收：
  A. 物調金額前一期修正了，本期「前期累計」還是舊數字（連帶總計也對不上）
  B. 計價總表把 30% 打成 30：累計至本期數量被誤植成固定值，不是公式
  C. 兩期同一欄的小數位數設定不同（0.00 vs 0.000000），數字一樣、畫面不一樣
  D. 計價總表沒有單價欄，「單價 x 數量 = 複價」這項會被略過（⚠）

openpyxl 存出來的公式沒有計算結果，核算腳本讀不到數字——產生後一定要用 Excel
開啟存檔一次（make_examples.ps1 會自動做）。
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).parent
AMT = '#,##0;-#,##0;"-"'
THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='E7EFE7')

# 請款明細表的工程項目：(項次, 名稱, 單位, 原契約數量, 單價, 第一期本期數量, 第二期本期數量)
ITEMS = [
    (1, '瀝青混凝土鋪面', 'm2', 2000, 350, 800, 700),
    (2, '預鑄混凝土路緣石', 'm', 800, 1200, 300, 250),
    (3, '人行道高壓磚', 'm2', 600, 1500, 200, 250),
    (4, '交通號誌設施', '式', 1, 250000, 0.3, 0.4),
]
PRICE_ADJ_P1 = 12000        # 第一期物調（已修正後的正確值）
PRICE_ADJ_P1_STALE = 12500  # 第一期物調（修正前的舊值）—— 錯誤 A：第二期還抄這個
PRICE_ADJ_P2 = 8000


def _header(ws, title, groups):
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = '※ 練習用假資料，工程名稱與金額均為虛構'
    ws['A2'].font = Font(color='C00000', size=10)
    ws['A4'], ws['B4'] = '項次', '工程項目'
    ws['A5'], ws['B5'] = '', ''
    for group, first_col, subs in groups:
        ws.cell(row=4, column=first_col, value=group)
        ws.merge_cells(start_row=4, start_column=first_col, end_row=4, end_column=first_col + len(subs) - 1)
        for i, s in enumerate(subs):
            ws.cell(row=5, column=first_col + i, value=s)
    last_col = max(fc + len(s) - 1 for _, fc, s in groups)
    for r in (4, 5):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border = HEAD_FILL, BOX
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions['A'].width = 7
    ws.column_dimensions['B'].width = 20
    for c in range(3, last_col + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 12
    return last_col


def _box_rows(ws, first, last, last_col):
    for r in range(first, last + 1):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).border = BOX


def detail_sheet(wb, period):
    """請款明細表：逐項明細，有單價。A 欄項次、C~F 原契約、G~H 前期、I~J 本期、K~L 累計。"""
    ws = wb.create_sheet('請款明細表')
    last_col = _header(ws, '示範道路改善工程　第{0}期估驗計價　請款明細表'.format('一' if period == 1 else '二'), [
        ('原契約', 3, ['單位', '數量', '單價', '複價']),
        ('前期累計', 7, ['數量', '複價']),
        ('本期完成', 9, ['數量', '複價']),
        ('累計至本期', 11, ['數量', '複價']),
    ])
    ws['A6'], ws['B6'] = '壹', '發包工程費'
    for col in 'FHJL':
        ws['{0}6'.format(col)] = '=SUM({0}7:{0}10)'.format(col)
    for i, (code, name, unit, qty0, price, q1, q2) in enumerate(ITEMS):
        r = 7 + i
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=unit)
        ws.cell(row=r, column=4, value=qty0)
        ws.cell(row=r, column=5, value=price)
        ws['F{0}'.format(r)] = '=ROUND(D{0}*E{0},0)'.format(r)
        # 前期累計：第一期是 0；第二期照實務做法，是承辦從第一期「抄」過來打進去的數字
        prev_q = 0 if period == 1 else q1
        ws['G{0}'.format(r)] = prev_q
        ws['H{0}'.format(r)] = 0 if period == 1 else round(price * q1)
        ws['I{0}'.format(r)] = q1 if period == 1 else q2
        ws['J{0}'.format(r)] = '=ROUND(E{0}*I{0},0)'.format(r)
        ws['K{0}'.format(r)] = '=G{0}+I{0}'.format(r)
        ws['L{0}'.format(r)] = '=H{0}+J{0}'.format(r)
    ws['B11'] = '物價調整金額'
    ws['H11'] = 0 if period == 1 else PRICE_ADJ_P1_STALE   # 錯誤 A
    ws['J11'] = PRICE_ADJ_P1 if period == 1 else PRICE_ADJ_P2
    ws['L11'] = '=H11+J11'
    ws['B12'] = '總計'
    for col in 'FHJL':
        ws['{0}12'.format(col)] = '={0}6+{0}11'.format(col) if col != 'F' else '=F6'
    for r in range(6, 13):
        for col in 'DGIK':
            ws['{0}{1}'.format(col, r)].number_format = '#,##0.00'
        for col in 'EFHJL':
            ws['{0}{1}'.format(col, r)].number_format = AMT
    _box_rows(ws, 6, 12, last_col)
    for r in (6, 12):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).font = Font(bold=True)
    return ws


def summary_sheet(wb, period):
    """計價總表：分類彙總，沒有單價欄（錯誤 D）。C~E 原契約、F~G 前期、H~I 本期、J~K 累計。"""
    ws = wb.create_sheet('計價總表')
    last_col = _header(ws, '示範道路改善工程　第{0}期估驗計價　計價總表'.format('一' if period == 1 else '二'), [
        ('原契約', 3, ['單位', '數量', '複價']),
        ('前期累計', 6, ['數量', '複價']),
        ('本期完成', 8, ['數量', '複價']),
        ('累計至本期', 10, ['數量', '複價']),
    ])
    d = "'請款明細表'!"
    ws['A6'], ws['B6'] = '壹', '發包工程費'
    for col in 'EGIK':
        ws['{0}6'.format(col)] = '=SUM({0}7:{0}8)'.format(col)
    rows = [
        # (列, 項次, 名稱, 明細表的列範圍)
        (7, '壹.一', '道路工程費', (7, 9)),
        (8, '壹.二', '交通工程費', (10, 10)),
    ]
    for r, code, name, (a, b) in rows:
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value='式')
        ws.cell(row=r, column=4, value=1)
        ws['E{0}'.format(r)] = '=SUM({0}F{1}:F{2})'.format(d, a, b)
        ws['G{0}'.format(r)] = '=SUM({0}H{1}:H{2})'.format(d, a, b)
        ws['I{0}'.format(r)] = '=SUM({0}J{1}:J{2})'.format(d, a, b)
        ws['F{0}'.format(r)] = '=G{0}/E{0}'.format(r)
        ws['H{0}'.format(r)] = '=I{0}/E{0}'.format(r)
        ws['J{0}'.format(r)] = '=F{0}+H{0}'.format(r)
        ws['K{0}'.format(r)] = '=G{0}+I{0}'.format(r)
    if period == 1:
        ws['J8'] = 30   # 錯誤 B：把 30% 打成 30，公式被固定值蓋掉
    ws['B9'] = '物價指數調整'
    ws['G9'], ws['I9'], ws['K9'] = '={0}H11'.format(d), '={0}J11'.format(d), '=G9+I9'
    ws['B10'] = '總計'
    ws['E10'] = '=E6'
    for col in 'GIK':
        ws['{0}10'.format(col)] = '={0}6+{0}9'.format(col)
    # 錯誤 C：兩期數量欄的小數位數設定不同
    qty_fmt = '0.00' if period == 1 else '0.000000'
    for r in range(6, 11):
        for col in 'FHJ':
            ws['{0}{1}'.format(col, r)].number_format = qty_fmt
        ws['D{0}'.format(r)].number_format = '0'
        for col in 'EGIK':
            ws['{0}{1}'.format(col, r)].number_format = AMT
    _box_rows(ws, 6, 10, last_col)
    for r in (6, 10):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).font = Font(bold=True)
    return ws


def build(period):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    detail_sheet(wb, period)
    summary_sheet(wb, period)
    name = '練習用-第{0}期估驗計價.xlsx'.format('一' if period == 1 else '二')
    wb.save(OUT / name)
    return OUT / name


if __name__ == '__main__':
    for p in (1, 2):
        print('已產生：{0}'.format(build(p)))
